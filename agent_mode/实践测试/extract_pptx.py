# -*- coding: utf-8 -*-
"""Extract full text content (incl. tables) from pptx files."""
import sys
from pptx import Presentation
from pptx.util import Emu

def iter_shapes(shapes, depth=0):
    for sh in shapes:
        yield sh, depth

def walk(shapes, out, depth=0):
    for sh in shapes:
        prefix = "  " * depth
        if sh.shape_type == 6:  # group
            out.append(f"{prefix}[GROUP]")
            walk(sh.shapes, out, depth + 1)
        elif sh.has_text_frame:
            txt = "\n".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            if txt.strip():
                out.append(f"{prefix}{txt}")
        if sh.has_table:
            tbl = sh.table
            out.append(f"{prefix}[TABLE {len(tbl.rows)}x{len(tbl.columns)}]")
            for r in tbl.rows:
                cells = [c.text.replace("\n", " ") for c in r.cells]
                out.append(f"{prefix}| " + " | ".join(cells))

def main(path):
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n===== SLIDE {i} =====")
        walk(slide.shapes, out)
    text = "\n".join(out)
    with open(path + ".txt", "w", encoding="utf-8") as f:
        f.write(text)
    print(f"Slides: {len(prs.slides)}")
    print(text[:4000])

if __name__ == "__main__":
    main(sys.argv[1])
