# -*- coding: utf-8 -*-
import openpyxl, os, io, sys

base = r"E:\教学\2026-2027-1\信息安全测评与风险评估\张玉\各类表格"
out = io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'tables_dump.txt'), 'w', encoding='utf-8')
for fn in sorted(os.listdir(base)):
    if not fn.endswith('.xlsx'): continue
    wb = openpyxl.load_workbook(os.path.join(base, fn), data_only=True)
    out.write('='*70 + '\n文件：' + fn + '\n' + '='*70 + '\n')
    for ws in wb.worksheets:
        out.write('\n--- Sheet: %s (rows=%d cols=%d) ---\n' % (ws.title, ws.max_row, ws.max_column))
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=True):
            cells = []
            for v in row:
                if v is None: cells.append('')
                else: cells.append(str(v).replace('\n', '⏎'))
            out.write(' | '.join(cells) + '\n')
out.close()
print('done')
